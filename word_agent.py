from pathlib import Path
from docx import Document as WordDocument
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


def WordReaderAgent(file_path):
    """Reads DOCX files and prepares content for RAG."""

    path = Path(file_path)

    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    if path.suffix.lower() != ".docx":
        raise ValueError("Please provide a DOCX file.")

    word_file = WordDocument(str(path))
    documents = []

    # Read paragraphs
    for index, paragraph in enumerate(word_file.paragraphs, start=1):
        if paragraph.text.strip():
            documents.append(
                Document(
                    page_content=paragraph.text,
                    metadata={
                        "source": path.name,
                        "paragraph": index,
                        "file_type": "DOCX"
                    }
                )
            )

    # Read tables
    for table_index, table in enumerate(word_file.tables, start=1):
        for row_index, row in enumerate(table.rows, start=1):
            row_text = " | ".join(
                cell.text.strip() for cell in row.cells
            )

            if row_text.strip():
                documents.append(
                    Document(
                        page_content=row_text,
                        metadata={
                            "source": path.name,
                            "table": table_index,
                            "row": row_index,
                            "file_type": "DOCX"
                        }
                    )
                )

    if not documents:
        raise ValueError("No readable text found in the DOCX file.")

    splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=150
    )

    chunks = splitter.split_documents(documents)

    return {
        "file_name": path.name,
        "total_paragraphs": len(word_file.paragraphs),
        "total_chunks": len(chunks),
        "text": "\n\n".join(
            document.page_content for document in documents
        ),
        "documents": chunks
    }